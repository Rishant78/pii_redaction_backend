import json
from pathlib import Path
from unittest.mock import patch
from src.evaluation import evaluate_rhp
from src.models import PIIType, Span
from docx import Document

def test_evaluate_rhp_strict_matching(tmp_path):
    doc_path = tmp_path / "test.docx"
    doc = Document()
    doc.add_paragraph("Kushal Hegde is at 123 Pune.")
    doc.save(doc_path)
    
    gt_path = tmp_path / "gt.json"
    
    def run_eval(spans, gt_annotations):
        with open(gt_path, 'w') as f:
            json.dump([{"paragraph_index": 0, "annotations": gt_annotations}], f)
            
        with patch('src.evaluation.detect_all', return_value=spans), \
             patch('src.document.discover_entities', return_value={}):
             return evaluate_rhp(doc_path, gt_path)

    # 1. Exact match -> 1 TP
    res = run_eval(
        [Span(0, 12, "Kushal Hegde", PIIType.PERSON, 1.0, "mock")],
        [{"start": 0, "end": 12, "type": "PERSON"}]
    )
    assert res['overall']['tp'] == 1
    assert res['overall']['fp'] == 0
    assert res['overall']['fn'] == 0
    assert res['per_type']['PERSON']['precision'] == 1.0
    
    # 2. Wrong boundary -> FP + FN
    res = run_eval(
        [Span(0, 10, "Kushal Heg", PIIType.PERSON, 1.0, "mock")],
        [{"start": 0, "end": 12, "type": "PERSON"}]
    )
    assert res['overall']['tp'] == 0
    assert res['overall']['fp'] == 1
    assert res['overall']['fn'] == 1

    # 3. Wrong type -> FP + FN
    res = run_eval(
        [Span(0, 12, "Kushal Hegde", PIIType.ORGANIZATION, 1.0, "mock")],
        [{"start": 0, "end": 12, "type": "PERSON"}]
    )
    assert res['overall']['tp'] == 0
    assert res['overall']['fp'] == 1
    assert res['overall']['fn'] == 1
    
    # 4. Missing gold -> FN
    res = run_eval(
        [],
        [{"start": 0, "end": 12, "type": "PERSON"}]
    )
    assert res['overall']['tp'] == 0
    assert res['overall']['fp'] == 0
    assert res['overall']['fn'] == 1
    
    # 5. Extra prediction -> FP
    res = run_eval(
        [Span(0, 12, "Kushal Hegde", PIIType.PERSON, 1.0, "mock")],
        []
    )
    assert res['overall']['tp'] == 0
    assert res['overall']['fp'] == 1
    assert res['overall']['fn'] == 0

def test_evaluate_rhp_zero_denominator(tmp_path):
    doc_path = tmp_path / "test.docx"
    doc = Document()
    doc.add_paragraph("No PII here.")
    doc.save(doc_path)
    
    gt_path = tmp_path / "gt.json"
    with open(gt_path, 'w') as f:
        json.dump([{"paragraph_index": 0, "annotations": []}], f)
        
    with patch('src.evaluation.detect_all', return_value=[]), \
         patch('src.document.discover_entities', return_value={}):
         res = evaluate_rhp(doc_path, gt_path)
         
    # Should not divide by zero
    assert res['overall']['precision'] == 0.0
    assert res['overall']['recall'] == 0.0
    assert res['overall']['f1'] == 0.0
