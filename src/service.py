from __future__ import annotations
from pathlib import Path
from .document import redact_docx, iter_paragraphs
from .detectors import detect_all
from .models import PIIType
from docx import Document


def analyze_docx(path: str|Path):
    doc=Document(str(path))
    detections=[]
    for p in iter_paragraphs(doc):
        if p.text:
            for s in detect_all(p.text):
                detections.append(s)
    counts={t.value:0 for t in PIIType}
    uniques={t.value:set() for t in PIIType}
    for s in detections:
        counts[s.pii_type.value]+=1
        uniques[s.pii_type.value].add(s.text)
    return {
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "detections": len(detections),
        "counts": counts,
        "unique_values": {k:len(v) for k,v in uniques.items()},
        "samples": [{"type":s.pii_type.value,"text":s.text,"confidence":s.confidence,"source":s.source} for s in detections[:100]],
    }
