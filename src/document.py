from __future__ import annotations
from pathlib import Path
from docx import Document
from docx.document import Document as DocumentType
from docx.text.paragraph import Paragraph
from docx.table import Table, _Cell
from .detectors import detect_all
from .models import Span, PIIType
from .replacements import ReplacementGenerator

def generate_person_aliases(full_name: str) -> list[str]:
    """Generate shortened aliases for a discovered person name."""
    words = full_name.split()
    aliases = []
    if len(words) >= 3:
        if words[0] and words[0][0].isupper() and words[-1] and words[-1][0].isupper():
            aliases.append(f"{words[0]} {words[-1]}")
    return aliases

def generate_org_aliases(full_name: str) -> list[str]:
    """Generate aliases for a discovered organization (e.g. without corporate suffix)."""
    aliases = []
    import re
    normalized = re.sub(r'\s+', ' ', full_name).strip()
    if normalized != full_name:
        aliases.append(normalized)
    
    suffix_pattern = re.compile(r'(?:\s+(?:Limited|Ltd\.?|Private Limited|Pvt\.?\s+Ltd\.?|LLP|L\.L\.P\.|Corporation|Inc\.?|Incorporated|Holdings|Industries))+$', re.I)
    without_suffix = suffix_pattern.sub('', normalized)
    if without_suffix != normalized and len(without_suffix.split()) >= 2:
        aliases.append(without_suffix)
    return aliases


def iter_paragraphs(parent):
    if isinstance(parent, DocumentType):
        for p in parent.paragraphs: yield p
        for t in parent.tables: yield from iter_table(t)
        for section in parent.sections:
            for p in section.header.paragraphs: yield p
            for t in section.header.tables: yield from iter_table(t)
            for p in section.footer.paragraphs: yield p
            for t in section.footer.tables: yield from iter_table(t)
    elif isinstance(parent, _Cell):
        for p in parent.paragraphs: yield p
        for t in parent.tables: yield from iter_table(t)


def iter_table(table: Table):
    for row in table.rows:
        for cell in row.cells:
            yield from iter_paragraphs(cell)


def discover_entities(doc: Document):  # type: ignore
    """Discover high-confidence entities first, then apply them consistently."""
    registry: dict[tuple[str,str], str] = {}
    ambiguous_aliases = set()

    for p in iter_paragraphs(doc):
        text=p.text or ""
        for s in detect_all(text):
            if s.pii_type.value in {"EMAIL","PHONE","SSN","CREDIT_CARD","IP_ADDRESS","DOB","ADDRESS","ORGANIZATION"} or s.confidence >= 0.90:
                key = (s.pii_type.value, s.text)
                
                # Normalize canonical value for organizations to handle case/whitespace variants
                if s.pii_type == PIIType.ORGANIZATION:
                    import re
                    canonical = re.sub(r'\s+', ' ', s.text).strip().title()
                else:
                    canonical = s.text
                    
                if key not in registry:
                    registry[key] = canonical

    new_aliases = {}
    for (type_name, original) in list(registry.keys()):
        canonical = registry[(type_name, original)]
        if type_name == PIIType.PERSON.value:
            for alias in generate_person_aliases(original):
                akey = (type_name, alias)
                if akey in registry:
                    continue
                if akey in new_aliases:
                    if new_aliases[akey] != original:
                        ambiguous_aliases.add(akey)
                else:
                    new_aliases[akey] = original
        elif type_name == PIIType.ORGANIZATION.value:
            for alias in generate_org_aliases(canonical):
                akey = (type_name, alias)
                if akey in registry:
                    continue
                if akey in new_aliases:
                    if new_aliases[akey] != canonical:
                        ambiguous_aliases.add(akey)
                else:
                    new_aliases[akey] = canonical
                    
    for akey, canonical in new_aliases.items():
        if akey not in ambiguous_aliases:
            registry[akey] = canonical

    return registry


def _replace_spans_in_paragraph(paragraph: Paragraph, spans: list[Span], generator: ReplacementGenerator):
    runs=paragraph.runs
    if not runs or not spans: return
    boundaries=[]; pos=0
    for idx,r in enumerate(runs):
        txt=r.text or ""
        boundaries.append((pos,pos+len(txt),idx)); pos+=len(txt)
    for span in sorted(spans,key=lambda x:x.start,reverse=True):
        affected=[b for b in boundaries if not (span.end<=b[0] or span.start>=b[1])]
        if not affected: continue
        first=affected[0][2]
        lookup_val = span.canonical_text if span.canonical_text else span.text
        replacement=generator.generate(lookup_val,span.pii_type)
        for start,end,idx in affected:
            txt=runs[idx].text or ""
            ls=max(span.start,start)-start; le=min(span.end,end)-start
            if idx==first:
                runs[idx].text=txt[:ls]+replacement+txt[le:]
            else:
                runs[idx].text=txt[:ls]+txt[le:]


def redact_docx(input_path: str|Path, output_path: str|Path):
    doc=Document(str(input_path))
    generator=ReplacementGenerator()
    registry=discover_entities(doc)
    # Map every unique discovered value to one replacement.
    for (type_name, _), canonical in registry.items():
        from .models import PIIType
        generator.generate(canonical, PIIType(type_name))

    all_unique=[]; seen=set()
    for p in iter_paragraphs(doc):
        text=p.text or ""
        if not text: continue
        spans=[]
        # First use current detections; then exact-match all high-confidence entities
        # discovered anywhere in the document so repeated mentions remain consistent.
        spans.extend(detect_all(text))
        import re
        for (type_name, lookup_string), canonical in registry.items():
            if type_name in {PIIType.PERSON.value, PIIType.ORGANIZATION.value}:
                pattern_parts = [re.escape(w) for w in lookup_string.split()]
                if not pattern_parts: continue
                if type_name == PIIType.PERSON.value:
                    pattern = r'\b' + r'\s+'.join(pattern_parts) + r'\b'
                else:
                    pattern = r'(?<![A-Za-z0-9])' + r'\s+'.join(pattern_parts) + r'(?![A-Za-z0-9])'
            else:
                pattern = re.escape(lookup_string)
            for match in re.finditer(pattern, text, re.I):
                from .models import PIIType
                spans.append(Span(match.start(),match.end(),match.group(),PIIType(type_name),1.0,"entity-registry", canonical_text=canonical))
        # resolve overlaps locally
        from .detectors import merge_and_resolve
        spans=merge_and_resolve(spans)
        _replace_spans_in_paragraph(p,spans,generator)
        for s in spans:
            key=(s.pii_type.value,s.text)
            if key not in seen:
                seen.add(key); all_unique.append(s)
    doc.save(str(output_path))
    return all_unique, generator
